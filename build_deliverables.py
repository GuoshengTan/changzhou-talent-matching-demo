from pathlib import Path
import json,html
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
B=Path(__file__).resolve().parent
D=json.loads((B/'public_data.json').read_text(encoding='utf8'));R=json.loads((B/'results.json').read_text(encoding='utf8'));C={c['id']:c for c in D['candidates']}
doc=Document();s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=Inches(.7);s.left_margin=s.right_margin=Inches(.8)
for nm in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
 st=doc.styles[nm];st.font.name='Microsoft YaHei';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(11 if nm=='Normal' else 25 if nm=='Title' else 17 if nm=='Heading 1' else 13)
 rf=st._element.get_or_add_rPr().get_or_add_rFonts()
 for k in list(rf.attrib):
  if 'theme' in k.lower():del rf.attrib[k]
 for k in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+k),'Microsoft YaHei')
 st.paragraph_format.space_after=Pt(8)
doc.styles['Normal'].paragraph_format.line_spacing=1.2
for nm in ['Title','Subtitle']:
 pp=doc.styles[nm]._element.find(qn('w:pPr'))
 if pp is not None:
  for x in list(pp):
   if x.tag==qn('w:pBdr'):pp.remove(x)
def p(t):return doc.add_paragraph(t)
def h(t):doc.add_heading(t,1)
def page():doc.add_page_break()
def tab(head,rows,widths):
 t=doc.add_table(rows=1,cols=len(head));t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for cell,val in zip(t.rows[0].cells,head):cell.text=val
 for row in rows:
  for cell,val in zip(t.add_row().cells,row):cell.text=str(val)
 for ri,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for ci,cell in enumerate(row.cells):
   cell.width=Inches(widths[ci]);pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'173343' if ri==0 else 'EEF3F5' if ri%2 else 'FFFFFF');pr.append(sh)
   borders=OxmlElement('w:tcBorders')
   for name in ['top','left','bottom','right']:
    el=OxmlElement('w:'+name);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   pr.append(borders)
   for pp in cell.paragraphs:
    pp.paragraph_format.space_after=Pt(6);pp.paragraph_format.space_before=Pt(4)
    for r in pp.runs:
     r.font.size=Pt(10)
     if ri==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
 return t
def link(text,url):
 from docx.opc.constants import RELATIONSHIP_TYPE as RT
 pp=p('');e=OxmlElement('w:hyperlink');e.set(qn('r:id'),pp.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'245A76');pr.append(co);r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);e.append(r);pp._p.append(e)
foot=s.footer.paragraphs[0];foot.text='公开资料匹配演示  ·  2026年9月29日  ·  ';fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
doc.add_heading('常州产业人才匹配\n公开数据演示报告',0)
p('第二项成果配套材料  真实公开资料与可复现原型')
p('本次已将原有虚构人才演示替换为10名科研人员的官方公开资料，结合常州2份需求材料抽取3个技术检索方向，实际完成30组需求与人才资源配对计算。结果用于展示研究能力检索、来源追溯及解释功能，尚不构成招聘推荐或模型效果验证。')
h('数据规模与来源')
tab(['组成','数量','口径'],[['需求材料','2份','常州大学招聘公告、常州市科技局需求报道'],['检索方向','3个','固态电解质、复材加工成型、机器人研发'],['人才资料','10人','高校及科研院所官方公开主页'],['配对计算','30组','3个方向×10名科研人员']],[1.1,.7,4.9])
h('本次实际使用的方法')
p('先将原文涉及的研究方向人工编码为技能标签，再用标签覆盖率和文本相似度进行融合排序。默认分数为0.7×标签覆盖率＋0.3×文本相似度。文字匹配使用字符二元组TF-IDF余弦相似度，词表及逆文档频率只由10份人才摘要建立。权重为演示配置，没有训练或优化。')
p('标签保留逐项来源释义。摘要由人工整理，不是完整简历，也未经过独立专家标注。程序没有启用BERT或协同过滤；公开网页不提供历史投递、点击或录用交互，不能据此补造协同过滤数据。')
p('附件包含可离线打开的交互演示页、JSON、CSV及Python脚本。后文说明样本选择、实际排序结果及可用于论文的表述。')
for start in [0,5]:
 page();h('公开科研人才资料 '+('一' if start==0 else '二'))
 p('资料检索于2026年9月29日。以下是科研资源样本；未核实求职意愿或合作意愿，未采集联系方式、出生日期。页面日期较早的资料须再次确认现状。')
 for c in D['candidates'][start:start+5]:
  doc.add_heading(c['id']+'  '+c['name']+'  '+c['organization'],2)
  p(c['summary']);p('编码标签：'+'、'.join(c['skills'])+'。来源：['+c['id']+']。')
 p('样本选择采用目的性检索，围绕三个演示方向纳入官方主页信息较完整者。样本规模小、机构集中，不能代表全国人才供给，也不能用跨领域样本的排序差异证明算法泛化能力。')
page();h('需求方向及实际匹配结果')
for q in D['queries']:
 doc.add_heading(q['id']+'  '+q['title'],2);p('来源：['+q['source']+']。'+q['scope']+'。查询标签：'+'、'.join(q['skills'])+'。')
 rows=R['rankings'][q['id']][:3]
 tab(['科研资源','演示分数','命中标签'],[[r['name'],f"{r['score']:.3f}",'、'.join(r['covered']) or '无'] for r in rows],[1.3,1,4.4])
page();h('结果解释与资格核验')
p('固态电解质方向的前列结果具有公开的电解质材料研究依据；复合材料方向优先返回涉及复材成型与碳纤维的研究人员。机器人方向原文只披露宽泛的研发和算法需求，三名机器人研究人员均可命中唯一需求标签，文本分数只能提供弱区分。此时应由需求单位补充技术任务，不能人为加上视觉、导航等要求来制造精准匹配。')
h('招聘条件不能被能力分数替代')
p('D01是博士后招聘，包含博士学位、博士毕业年限、年龄、论文及英语能力等条件。此次公开教授和研究人员样本并非应聘者，不应因为研究方向接近就标为合格博士后。原型统一将应聘资格设为未核验，仅把对应技术方向用于科研资源检索。若用于正式招聘，需另建近期博士样本并逐条核对硬条件。')
h('权重敏感性检查')
tab(['方向','标签权重0.5','标签权重0.7','标签权重0.9'],[[q['id']]+['、'.join(C[c]['name'] for c in R['sensitivity_top3'][str(a)][q['id']]) for a in [.5,.7,.9]] for q in D['queries']],[.55,2.05,2.05,2.05])
p('上表为同一小样本下前三位排序随权重变化的检查，不是准确率评价。由于缺少独立相关性标签，未计算Precision、Recall、NDCG或录用率。')
h('可写入论文的事实')
p('基于2份常州公开需求材料和10份科研人员官方主页，研究整理了3个技术检索方向与30组配对记录，以显式标签覆盖和文本相似度实现可解释检索原型，并展示来源追溯和权重敏感性。该案例验证了程序流程可运行，尚不能证明推荐效果优于其他方法。')
page();h('来源索引与复现步骤')
for src in D['sources']:
 doc.add_heading(src['id']+'  '+src['title'],2);p('发布日期：'+src['date']+'。'+src['summary']);p('原文条件：'+src['constraints']);link('打开需求原文',src['url'])
h('运行演示')
p('双击“打开演示.html”，选择技术方向，调整标签权重，查看命中标签、来源和未核验事项。页面不联网计算，不自动联系任何人。')
p('数据已存放在public_data.json。使用Python 3运行：python match_public.py。程序仅依赖标准库，重新生成results.json和全部匹配结果.csv。修改数据后先运行匹配脚本，再运行build_deliverables.py更新展示页与Word报告；后者需python-docx。')
p('人才与需求CSV供人工审阅。规范化标签、释义和链接保留在JSON中。需求Q01与Q02属于同一招聘材料，不能重复计算招聘主体或人数。网页源文可能更新，复核应同时记录检索日期；本包保留人工结构化摘要，不作为网页全文存档。')
page();h('人才资料来源索引')
for c in D['candidates']:
 link('['+c['id']+'] '+c['name']+' · '+c['organization']+'官方主页',c['url'])
p('全部来源在2026年9月29日检索。部分页面通过搜索引擎索引读取；当直连页面无法解析时，保留官方域名索引内容及原文链接作为本次资料依据。实际对接前应由人员复核最新页面。')
h('下一步正式验证所需数据')
p('针对同一个明确技术任务，补充同领域且资历相近的人才池；让至少两名领域评审独立判断适配程度，再处理分歧。评审时不展示算法分数，保留评审依据。人才合作意愿与企业反馈应单独记录，不能从学术主页推断。')
p('若用于博士后招聘，需要符合该招聘群体的近期博士数据及资格证据；若用于校企合作，应使用企业确认的合作任务。真实交互数据取得后，才具备增加协同过滤和开展独立效果评价的条件。')
doc.core_properties.author=''
doc.core_properties.comments=''
doc.core_properties.last_modified_by=''
doc.save(B/'真实公开数据与匹配演示报告.docx')
# Offline interactive viewer reweights actual components calculated in Python.
template='''<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>常州产业人才匹配演示</title><style>
*{box-sizing:border-box}body{margin:0;background:#f2f5f4;color:#19323d;font:16px/1.65 "Microsoft YaHei",sans-serif}main{max-width:1180px;margin:40px auto;padding:0 26px}header{border-bottom:1px solid #cad7d5;padding-bottom:24px}h1{font-size:34px;margin:10px 0}h2{font-size:22px}small,.muted{color:#62757b}.stats{display:flex;gap:16px;margin:25px 0}.stat{background:white;padding:15px 25px;border-radius:10px;flex:1}.stat b{font-size:28px;color:#226c64;display:block}.controls{display:flex;gap:30px;align-items:center;background:#173c42;color:white;padding:22px;border-radius:12px;flex-wrap:wrap}select{font:inherit;padding:10px;border-radius:6px;max-width:100%}input{accent-color:#72c3af}.note{background:#e4efea;padding:18px;border-radius:10px;margin:20px 0}article{background:white;padding:25px;margin:16px 0;border-radius:12px;border:1px solid #dbe4e1}.head{display:flex;justify-content:space-between;gap:20px}.score{font-size:26px;color:#23695f}.tag{display:inline-block;background:#e1eee8;margin:3px;padding:3px 10px;border-radius:5px}a{color:#276978}details{border-top:1px solid #ddd;padding-top:10px;margin-top:12px}button{font:inherit;border:0;padding:8px 18px;background:#d9ebe5;color:#173c42;border-radius:5px;cursor:pointer}@media(max-width:650px){.stats{flex-wrap:wrap}.stat{flex-basis:40%}h1{font-size:26px}}
</style><main><header><small>CHANGZHOU / PUBLIC RESEARCH DATA</small><h1>产业技术需求与科研人才资源匹配</h1><div class="muted">真实公开资料演示 · 2026年9月29日 · 可追溯来源</div></header><div class="stats"><div class="stat"><b>10</b>科研人员主页</div><div class="stat"><b>2</b>需求来源材料</div><div class="stat"><b>3</b>技术检索方向</div><div class="stat"><b>30</b>配对计算</div></div><div class="controls"><label>技术方向 <select id="query"></select></label><label>标签权重 <input id="alpha" type="range" min="0" max="100" value="70"> <b id="weight">0.70</b></label><button id="export">导出当前结果</button></div><div class="note" id="context"></div><p class="muted">分数 = 标签覆盖率 × 权重 + 文本相似度 × 剩余权重。仅排序科研方向，不表示求职意愿、招聘资格或成功率。样本为目的性选择，尚无独立准确率评价。</p><div id="results"></div><footer class="muted">原型使用显式标签和字符二元组TF-IDF；未启用BERT和协同过滤。零分表示当前摘要未提供对应证据。全部计算在本地完成。</footer></main><script>
const D=__DATA__,R=__RESULTS__;const esc=s=>String(s).replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));let current=[];
const sel=document.querySelector('#query'),slider=document.querySelector('#alpha');D.queries.forEach(q=>{const o=document.createElement('option');o.value=q.id;o.textContent=q.title;sel.append(o)});
function render(){const q=D.queries.find(q=>q.id===sel.value),source=D.sources.find(s=>s.id===q.source),a=Number(slider.value)/100;document.querySelector('#weight').textContent=a.toFixed(2);document.querySelector('#context').innerHTML=esc(q.scope)+'。<br>需求标签：'+q.skills.map(esc).join('、')+'。 <a target="_blank" rel="noopener" href="'+esc(source.url)+'">查看需求来源 '+q.source+'</a><br><small>'+esc(source.constraints)+'</small>';current=R.rankings[q.id].map(r=>({...r,score:a*r.coverage+(1-a)*r.text_similarity})).sort((x,y)=>y.score-x.score||x.candidate_id.localeCompare(y.candidate_id));document.querySelector('#results').innerHTML=current.map((r,i)=>{const c=D.candidates.find(c=>c.id===r.candidate_id);return `<article><div class="head"><div><h2>${i+1}. ${esc(c.name)} <small>${esc(c.organization)}</small></h2></div><div class="score">${r.score.toFixed(3)}<br><small>演示分数</small></div></div><p>${esc(c.summary)}</p><div>${r.covered.map(k=>'<span class="tag">'+esc(k)+'</span>').join('')||'<span class="muted">没有命中需求标签</span>'}</div><p class="muted">标签覆盖 ${r.coverage.toFixed(3)} · 文本相似度 ${r.text_similarity.toFixed(3)} · 尚无证据：${r.not_evidenced.map(esc).join('、')||'需求标签均有对应信息'}</p><details><summary>查看来源与标签依据</summary><p>${Object.entries(c.skill_evidence).map(([k,v])=>esc(k)+'：'+esc(v)).join('<br>')}</p><a target="_blank" rel="noopener" href="${esc(c.url)}">${esc(c.name)}官方资料</a><p class="muted">合作意愿未知；招聘资格未核验。上列证据为研究者释义，非原文引述。</p></details></article>`}).join('')}
sel.addEventListener('change',render);slider.addEventListener('input',render);document.querySelector('#export').onclick=()=>{const b=new Blob([JSON.stringify({query:sel.value,tag_weight:Number(slider.value)/100,results:current},null,2)],{type:'application/json'}),u=URL.createObjectURL(b),a=document.createElement('a');a.href=u;a.download='当前匹配结果.json';a.click();setTimeout(()=>URL.revokeObjectURL(u),1000)};render();
</script></html>'''
(B/'打开演示.html').write_text(template.replace('__DATA__',json.dumps(D,ensure_ascii=False)).replace('__RESULTS__',json.dumps(R,ensure_ascii=False)),encoding='utf8')
(B/'使用说明.txt').write_text('先双击 打开演示.html。选择需求方向、调整权重、展开来源依据，或导出当前结果。\n人才CSV与需求CSV均可用Excel编辑，但计算以public_data.json为准。\n复现：python match_public.py。仅依赖Python标准库。\nWord报告：真实公开数据与匹配演示报告.docx。\n结果范围：科研能力检索演示；不是求职者库或真实招聘推荐。\n数据：2份需求原文、3个查询方向、10名科研人员、30组配对。所有样本均为公开来源，无虚构简历。\n',encoding='utf8')
print('Word and HTML created')
